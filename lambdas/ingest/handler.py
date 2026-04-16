"""
Ingest Lambda for Financial Reports RAG.

Triggered automatically when a supported file is uploaded to S3 under the reports/ prefix.

Expected S3 key structure:
    reports/{TICKER}/{report_type}/{fiscal_period}/{filename}.{ext}

Supported formats:
    .pdf   — Annual/quarterly reports downloaded from SEC or investor relations pages
    .docx  — Word documents (analyst reports, internal notes)
    .xlsx  — Excel spreadsheets (financial tables, data sheets)
    .csv   — Comma-separated data files
    .html  — SEC EDGAR HTML filings (.htm also supported)
    .txt   — Plain text files

Examples:
    reports/AAPL/annual/2024/AAPL-10K-2024.pdf
    reports/NVDA/annual/2024/NVDA-10K-2024.pdf
    reports/AMZN/quarterly/2024Q3/AMZN-10Q-2024Q3.pdf
    reports/GOOGL/annual/2024/GOOGL-10K-2024.html

Process:
    1. Download file from S3
    2. Parse S3 key to extract financial metadata (ticker, report_type, fiscal_period)
    3. Extract full text using the appropriate extractor for the file format
    4. Split text into overlapping chunks (1000 chars, 200-char overlap)
    5. Embed each chunk with Amazon Titan Embed v2 (1024 dimensions)
    6. Store chunk + embedding + metadata in OpenSearch Serverless
"""

import json
import os
import hashlib
import boto3
import requests
from requests_aws4auth import AWS4Auth
from io import BytesIO

s3_client = boto3.client("s3")
bedrock   = boto3.client("bedrock-runtime", region_name=os.environ.get("BEDROCK_REGION", "us-east-1"))

OPENSEARCH_ENDPOINT = os.environ["OPENSEARCH_ENDPOINT"].rstrip("/")
INDEX_NAME          = os.environ.get("INDEX_NAME", "financial-reports")
REGION              = os.environ["AWS_REGION"]

CHUNK_SIZE      = 1000   # characters per chunk (larger than generic RAG — financial text is dense)
CHUNK_OVERLAP   = 200    # overlap so figures/sentences aren't cut off at boundaries
EMBEDDING_MODEL = "amazon.titan-embed-text-v2:0"
EMBEDDING_DIM   = 1024
MIN_CHUNK_LEN   = 100    # discard near-empty chunks


# ── OpenSearch helpers ─────────────────────────────────────────────────────────

def _get_awsauth() -> AWS4Auth:
    """Build AWS4Auth for signing OpenSearch Serverless requests using the Lambda execution role."""
    creds = boto3.Session().get_credentials().get_frozen_credentials()
    return AWS4Auth(creds.access_key, creds.secret_key, REGION, "aoss",
                    session_token=creds.token)


def _signed_request(method: str, url: str, body_bytes: bytes) -> tuple[int, bytes]:
    """Make a SigV4-signed HTTP request to OpenSearch Serverless. Returns (status, data)."""
    resp = requests.request(
        method, url,
        data=body_bytes,
        headers={"Content-Type": "application/json"},
        auth=_get_awsauth(),
    )
    return resp.status_code, resp.content


def opensearch_put(path: str, body: dict) -> dict:
    """PUT an index definition to OpenSearch Serverless."""
    url         = f"{OPENSEARCH_ENDPOINT}{path}"
    body_bytes  = json.dumps(body).encode()
    status, data = _signed_request("PUT", url, body_bytes)
    if status >= 400:
        raise RuntimeError(f"OpenSearch PUT {path} failed [{status}]: {data[:500].decode()}")
    return json.loads(data)


def opensearch_post(path: str, body: dict) -> dict:
    """POST a document to OpenSearch Serverless (auto-generates document ID)."""
    url         = f"{OPENSEARCH_ENDPOINT}{path}"
    body_bytes  = json.dumps(body).encode()
    status, data = _signed_request("POST", url, body_bytes)
    if status >= 400:
        raise RuntimeError(f"OpenSearch POST {path} failed [{status}]: {data[:500].decode()}")
    return json.loads(data)


# ── Index creation ─────────────────────────────────────────────────────────────

INDEX_MAPPING = {
    "settings": {
        "index": {
            "knn": True
        }
    },
    "mappings": {
        "properties": {
            "embedding":     {"type": "knn_vector", "dimension": EMBEDDING_DIM},
            "text":          {"type": "text"},
            "ticker":        {"type": "keyword"},
            "report_type":   {"type": "keyword"},
            "fiscal_period": {"type": "keyword"},
            "source":        {"type": "keyword"},
            "chunk_index":   {"type": "integer"},
            "doc_id":        {"type": "keyword"},
        }
    },
}


def create_index_if_not_exists():
    """
    Create the vector index with knn_vector mapping on first use.
    Safe to call on every invocation — silently skips if already exists.
    """
    url        = f"{OPENSEARCH_ENDPOINT}/{INDEX_NAME}"
    body_bytes = json.dumps(INDEX_MAPPING).encode()
    status, data = _signed_request("PUT", url, body_bytes)

    if status in (200, 201):
        print(f"Index '{INDEX_NAME}' created successfully.")
        return

    # Check if failure is "already exists" — that's fine
    try:
        err = json.loads(data).get("error", {})
        if "resource_already_exists_exception" in err.get("type", ""):
            print(f"Index '{INDEX_NAME}' already exists, skipping.")
            return
    except Exception:
        pass

    # Any other error is a real problem
    raise RuntimeError(f"Failed to create index '{INDEX_NAME}' [{status}]: {data[:500].decode()}")


# ── Text extraction ────────────────────────────────────────────────────────────

def extract_text_from_pdf(file_bytes: bytes) -> str:
    """Extract all text from a PDF file using pypdf. Pages joined with double newlines."""
    from pypdf import PdfReader
    reader = PdfReader(BytesIO(file_bytes))
    pages  = []
    for page in reader.pages:
        text = page.extract_text()
        if text and text.strip():
            pages.append(text.strip())
    full_text = "\n\n".join(pages)
    print(f"Extracted {len(full_text):,} characters from {len(reader.pages)} PDF pages")
    return full_text


def extract_text_from_docx(file_bytes: bytes) -> str:
    """Extract all text from a Word (.docx) file using python-docx."""
    import docx
    doc        = docx.Document(BytesIO(file_bytes))
    paragraphs = [p.text.strip() for p in doc.paragraphs if p.text.strip()]
    full_text  = "\n\n".join(paragraphs)
    print(f"Extracted {len(full_text):,} characters from {len(paragraphs)} paragraphs")
    return full_text


def extract_text_from_xlsx(file_bytes: bytes) -> str:
    """
    Extract all text from an Excel (.xlsx) file using openpyxl.
    Each sheet is rendered as tab-separated rows, sheets separated by double newlines.
    """
    import openpyxl
    wb     = openpyxl.load_workbook(BytesIO(file_bytes), read_only=True, data_only=True)
    sheets = []
    for sheet in wb.worksheets:
        rows = []
        for row in sheet.iter_rows(values_only=True):
            cells = [str(c) if c is not None else "" for c in row]
            row_text = "\t".join(cells).strip()
            if row_text:
                rows.append(row_text)
        if rows:
            sheets.append(f"[Sheet: {sheet.title}]\n" + "\n".join(rows))
    full_text = "\n\n".join(sheets)
    print(f"Extracted {len(full_text):,} characters from {len(wb.worksheets)} sheet(s)")
    return full_text


def extract_text_from_csv(file_bytes: bytes) -> str:
    """Extract all text from a CSV file. Rows rendered as tab-separated values."""
    import csv
    import io
    text   = file_bytes.decode("utf-8", errors="replace")
    reader = csv.reader(io.StringIO(text))
    rows   = ["\t".join(row) for row in reader if any(cell.strip() for cell in row)]
    full_text = "\n".join(rows)
    print(f"Extracted {len(full_text):,} characters from {len(rows)} CSV rows")
    return full_text


def extract_text_from_html(file_bytes: bytes) -> str:
    """
    Extract visible text from an HTML file using BeautifulSoup.
    Strips tags, scripts, and styles — useful for SEC EDGAR HTML filings.
    """
    from bs4 import BeautifulSoup
    soup    = BeautifulSoup(file_bytes, "html.parser")
    # Remove non-visible elements
    for tag in soup(["script", "style", "meta", "head"]):
        tag.decompose()
    full_text = soup.get_text(separator="\n")
    # Collapse excessive blank lines
    lines     = [line.strip() for line in full_text.splitlines() if line.strip()]
    full_text = "\n\n".join(lines)
    print(f"Extracted {len(full_text):,} characters from HTML")
    return full_text


def extract_text_from_txt(file_bytes: bytes) -> str:
    """Decode a plain text file."""
    full_text = file_bytes.decode("utf-8", errors="replace").strip()
    print(f"Extracted {len(full_text):,} characters from plain text")
    return full_text


# Format dispatcher
SUPPORTED_EXTENSIONS = {
    ".pdf":  extract_text_from_pdf,
    ".docx": extract_text_from_docx,
    ".xlsx": extract_text_from_xlsx,
    ".csv":  extract_text_from_csv,
    ".html": extract_text_from_html,
    ".htm":  extract_text_from_html,
    ".txt":  extract_text_from_txt,
}


def extract_text(file_bytes: bytes, key: str) -> str:
    """
    Dispatch to the correct extractor based on the file extension in the S3 key.
    Raises ValueError for unsupported formats.
    """
    ext      = "." + key.rsplit(".", 1)[-1].lower() if "." in key else ""
    extractor = SUPPORTED_EXTENSIONS.get(ext)
    if not extractor:
        supported = ", ".join(SUPPORTED_EXTENSIONS.keys())
        raise ValueError(
            f"Unsupported file format '{ext}' for key '{key}'. "
            f"Supported formats: {supported}"
        )
    return extractor(file_bytes)


# ── Text chunking ──────────────────────────────────────────────────────────────

def chunk_text(text: str) -> list[str]:
    """
    Split text into overlapping chunks.
    Overlap preserves context at chunk boundaries — important for financial figures
    that may span a sentence (e.g. "...revenue of $383.3 billion, up 8% year-over-year...").
    """
    chunks = []
    start  = 0
    while start < len(text):
        end   = min(start + CHUNK_SIZE, len(text))
        chunk = text[start:end].strip()
        if len(chunk) >= MIN_CHUNK_LEN:
            chunks.append(chunk)
        start += CHUNK_SIZE - CHUNK_OVERLAP
    return chunks


# ── Embeddings ─────────────────────────────────────────────────────────────────

def embed_text(text: str) -> list[float]:
    """
    Embed text using Amazon Titan Embed Text v2.
    Produces normalized 1024-dimensional vectors for cosine similarity search.
    """
    response = bedrock.invoke_model(
        modelId=EMBEDDING_MODEL,
        contentType="application/json",
        accept="application/json",
        body=json.dumps({
            "inputText": text[:8000],   # model's practical input limit
            "dimensions": EMBEDDING_DIM,
            "normalize": True,
        }),
    )
    return json.loads(response["body"].read())["embedding"]


# ── S3 key parsing ─────────────────────────────────────────────────────────────

def parse_s3_key(key: str) -> dict:
    """
    Parse the S3 key to extract financial metadata.

    Expected: reports/{ticker}/{report_type}/{fiscal_period}/{filename}
    Returns a dict with ticker, report_type, fiscal_period, filename.
    Falls back to UNKNOWN values if the key doesn't match the expected pattern.
    """
    parts = key.split("/")
    if len(parts) >= 5 and parts[0] == "reports":
        return {
            "ticker":        parts[1].upper(),
            "report_type":   parts[2].lower(),
            "fiscal_period": parts[3],
            "filename":      "/".join(parts[4:]),
        }
    # Flat upload fallback
    print(f"WARNING: Key '{key}' does not match expected pattern "
          f"reports/{{ticker}}/{{type}}/{{period}}/{{file}}. Using UNKNOWN metadata.")
    return {
        "ticker":        "UNKNOWN",
        "report_type":   "unknown",
        "fiscal_period": "unknown",
        "filename":      key.split("/")[-1],
    }


# ── Storage ────────────────────────────────────────────────────────────────────

def store_chunk(chunk: str, embedding: list[float], doc_id: str,
                chunk_index: int, source_key: str, metadata: dict):
    """Index a single chunk with its embedding and financial metadata in OpenSearch."""
    doc = {
        "text":          chunk,
        "embedding":     embedding,
        "source":        source_key,
        "chunk_index":   chunk_index,
        "doc_id":        doc_id,
        "ticker":        metadata["ticker"],
        "report_type":   metadata["report_type"],
        "fiscal_period": metadata["fiscal_period"],
    }
    # OpenSearch Serverless VECTORSEARCH does not support caller-specified document IDs.
    # Use POST to auto-generate the ID; doc_id + chunk_index are stored in the body for reference.
    opensearch_post(f"/{INDEX_NAME}/_doc", doc)


# ── Entry point ────────────────────────────────────────────────────────────────

def handler(event, context):
    """
    Lambda entry point — invoked by S3 for each file upload under reports/.
    Supports .pdf, .docx, .xlsx, .csv, .html/.htm, and .txt formats.
    One invocation may contain multiple S3 records (batch uploads).
    """
    create_index_if_not_exists()

    for record in event["Records"]:
        bucket = record["s3"]["bucket"]["name"]
        key    = record["s3"]["object"]["key"]

        ext = "." + key.rsplit(".", 1)[-1].lower() if "." in key else ""
        if ext not in SUPPORTED_EXTENSIONS:
            print(f"Skipping unsupported format '{ext}': {key}")
            continue

        print(f"\n--- Processing: s3://{bucket}/{key} ---")
        metadata = parse_s3_key(key)
        print(f"Metadata: ticker={metadata['ticker']} | "
              f"type={metadata['report_type']} | period={metadata['fiscal_period']}")

        # Download file
        file_bytes = s3_client.get_object(Bucket=bucket, Key=key)["Body"].read()
        print(f"Downloaded {len(file_bytes):,} bytes")

        # Extract text using the appropriate extractor for the file format
        text = extract_text(file_bytes, key)
        if not text.strip():
            print(f"WARNING: No text extracted from {key}. Skipping.")
            continue

        # Chunk
        chunks = chunk_text(text)
        print(f"Split into {len(chunks)} chunks")

        # Stable doc ID derived from the S3 key
        doc_id = hashlib.md5(key.encode()).hexdigest()[:12]

        # Embed and store each chunk
        for i, chunk in enumerate(chunks):
            embedding = embed_text(chunk)
            store_chunk(chunk, embedding, doc_id, i, key, metadata)
            if (i + 1) % 10 == 0 or (i + 1) == len(chunks):
                print(f"  Indexed {i + 1}/{len(chunks)} chunks")

        print(f"Done: {key} ({ext}) → {len(chunks)} chunks indexed")

    return {"statusCode": 200, "body": "Ingestion complete"}
